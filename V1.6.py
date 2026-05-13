import tkinter as tk
from tkinter import messagebox, simpledialog
import win32com.client as win32
import pythoncom
import os
import time
import pandas as pd
import re
import unicodedata
from datetime import datetime
import threading
import sys
import zipfile
import win32api
import win32con
import shutil
import glob
import subprocess
import webbrowser
import tempfile
from pywinauto import Application
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import customtkinter as ctk
from PIL import Image

# ================== CONFIGURAÇÕES ==================
TXT_ULTIMO_ENVIO = "ultimo_salvamento.txt"
TXT_ENVIO = "ultimos_envios.txt"
TXT_COBRANCA = "ultimas_cobrancas.txt"
TXT_PROCURACOES = "Procuracoes.txt"
TXT_ULTIMO_AJUSTE = "ultimo_ajuste.txt"
TXT_ULTIMO_GERAR = "ultimo_gerar.txt"

# Configurações padrão (usuário Victor)
BASE_DIR_VICTOR = r"C:\Users\TI Medicina\Desktop\ESOCIAL - PYTHON"
CAMINHO_CONTROLE_VICTOR = os.path.join(
    BASE_DIR_VICTOR, "Controle do E-Social.xlsx")
EMAIL_CONTA_PADRAO_VICTOR = "ti@medicinaoswaldocruz.com.br"
PASTA_PLANILHAS_VICTOR = os.path.join(BASE_DIR_VICTOR, "Planilhas")
PASTA_DOWNLOAD_VICTOR = r"C:\Users\TI Medicina\Downloads"

# Configurações Thailany
BASE_DIR_THAILANY = r"D:\Desktop\Oswaldo Cruz Gestão Ocupacional\E-Social"
CAMINHO_CONTROLE_THAILANY = os.path.join(
    BASE_DIR_THAILANY, "Controle do E-Social.xlsx")
EMAIL_CONTA_PADRAO_THAILANY = "esocial@medicinaoswaldocruz.com.br"
EMAIL_PLANILHAS_THAILANY = "planilhasesocial@outlook.com"
PASTA_PLANILHAS_THAILANY = os.path.join(
    BASE_DIR_THAILANY, "Planilhas Mensais - Conferência")
PASTA_DOWNLOAD_THAILANY = r"D:\Downloads"

CAMINHO_PATINHO_1 = os.path.join(BASE_DIR_VICTOR, "patinho.gif")
CAMINHO_PATINHO_2 = os.path.join(BASE_DIR_THAILANY, "patinho.gif")

# --- Novas definições de fonte ---
FONT_FAMILY = "Magsen-Regular"
# Fontes para CustomTkinter (usando tuplas ou nomes)
FONT_NORMAL = (FONT_FAMILY, 12)
FONT_BOLD = (FONT_FAMILY, 12, "bold")
FONT_ENTRY = (FONT_FAMILY, 12)
FONT_LABEL = (FONT_FAMILY, 13)
FONT_BUTTON = (FONT_FAMILY, 12, "bold")
FONT_FORGOT_PASSWORD = (FONT_FAMILY, 12, "underline")
FONT_TITLE = (FONT_FAMILY, 18, "bold")

MAX_LINHAS_COLUNA = 5
IDX_AK_LABEL = 36
IDX_AL_ANEXOS = 37
IDX_AM_ASSUNTO = 38
IDX_EMAIL_TO = 32
IDX_EMAIL_CC = 33
IDX_F_BLOQUEADA = 5  # Coluna F para verificar "BLOQUEADA"
IDX_AA_STATUS = 26  # Coluna AA para verificar "OK"
IDX_AC_VALIDADE = 28  # Coluna AC para validade da procuração
IDX_AE_PROCURACAO_TO = 30  # Coluna AE para email de procuração
IDX_AF_PROCURACAO_CC = 31  # Coluna AF para email de procuração em cópia
EXCLUIR_NA_COBRANCA = {"GRUPO OSWALDO CRUZ"}

# Configurações de Aparência Corporativa
ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue") # Tema azul corporativo

# ================== FUNÇÃO GERAR PENDENTES ==================

def executar_gerar_planilha(app_usuario, soc_usuario, soc_senha, soc_id, app_instance, btn_widget, status_banner):
    if app_usuario == "thailany":
        download_dir = PASTA_DOWNLOAD_THAILANY
        desktop_dir = os.path.dirname(os.path.dirname(BASE_DIR_THAILANY))
    else:
        download_dir = PASTA_DOWNLOAD_VICTOR
        desktop_dir = os.path.dirname(BASE_DIR_VICTOR)

    chrome_options = webdriver.ChromeOptions()
    prefs = {
        "download.default_directory": download_dir,
        "profile.default_content_setting_values.notifications": 2,
        "profile.default_content_setting_values.popups": 2,
        "download.prompt_for_download": False
    }
    chrome_options.add_experimental_option("prefs", prefs)
    chrome_options.add_argument("--start-maximized")

    temp_profile_dir = tempfile.mkdtemp(prefix="chrome_profile_eai_")
    chrome_options.add_argument(f"--user-data-dir={temp_profile_dir}")
    print(f"✨ Usando perfil de usuário temporário: {temp_profile_dir}")

    service = Service()
    driver = webdriver.Chrome(service=service, options=chrome_options)
    wait = WebDriverWait(driver, 20)

    try:
        pythoncom.CoInitialize()
        driver.get("https://sistema.soc.com.br/WebSoc/")

        wait.until(EC.presence_of_element_located(
            (By.NAME, "usu"))).send_keys(soc_usuario)
        driver.find_element(By.NAME, "senha").send_keys(soc_senha)

        teclado_div = wait.until(
            EC.presence_of_element_located((By.ID, "pteclado")))
        botoes = teclado_div.find_elements(By.TAG_NAME, "input")

        botoes_map = {btn.get_attribute("value"): btn for btn in botoes}

        for digito in soc_id:
            if digito not in botoes_map:
                raise ValueError(
                    f"Dígito {digito} não encontrado no teclado virtual.")
            botoes_map[digito].click()
            time.sleep(0.25)

        wait.until(EC.element_to_be_clickable((By.ID, "bt_entrar"))).click()

        campo_cod_prog = wait.until(
            EC.element_to_be_clickable(
                (By.ID, "cod_programa")))
        campo_cod_prog.click()
        campo_cod_prog.clear()
        campo_cod_prog.send_keys("267")

        wait.until(EC.element_to_be_clickable((By.ID, "btn_programa"))).click()

        excel_icon_found = False
        for iframe in driver.find_elements(By.TAG_NAME, "iframe"):
            try:
                driver.switch_to.frame(iframe)
                WebDriverWait(driver, 1).until(
                    EC.element_to_be_clickable(
                        (By.XPATH, "//img[contains(@src, 'excel.png')]"))
                ).click()
                excel_icon_found = True
                driver.switch_to.default_content()
                break
            except Exception:
                driver.switch_to.default_content()
        if not excel_icon_found:
            raise Exception("Ícone do Excel para gerar pedido não encontrado.")

        consultar_btn_found = False
        for iframe in driver.find_elements(By.TAG_NAME, "iframe"):
            try:
                driver.switch_to.frame(iframe)
                WebDriverWait(driver, 1).until(
                    EC.element_to_be_clickable(
                        (By.CSS_SELECTOR, "a.botaoT.btn-consultar-pedido"))
                ).click()
                consultar_btn_found = True
                driver.switch_to.default_content()
                break
            except Exception:
                driver.switch_to.default_content()
        if not consultar_btn_found:
            raise Exception("Botão 'Consultar Pedidos' não encontrado.")

        start_time = time.time()
        timeout = 300
        poll_interval = 3
        download_found = False
        print("🕒 Buscando o link de download (até 5 minutos)...")

        while time.time() - start_time < timeout:
            for iframe in driver.find_elements(By.TAG_NAME, "iframe"):
                try:
                    driver.switch_to.frame(iframe)
                    WebDriverWait(driver, 1).until(
                        EC.presence_of_element_located(
                            (By.CSS_SELECTOR, "div.div-download"))
                    )

                    maior_id = -1
                    div_maior = None
                    divs = driver.find_elements(
                        By.CSS_SELECTOR, "div.div-download")
                    for div in divs:
                        div_id = div.get_attribute("id")
                        match = re.search(r"(\d+)-download", div_id)
                        if match:
                            id_num = int(match.group(1))
                            if id_num > maior_id:
                                maior_id = id_num
                                div_maior = div

                    if div_maior:
                        div_maior.find_element(By.TAG_NAME, "a").click()
                        download_found = True

                    driver.switch_to.default_content()
                    if download_found:
                        break
                except Exception:
                    driver.switch_to.default_content()

            if download_found:
                break

            print(
                f"🕒 Link não encontrado. Clicando em 'Procurar' e aguardando {poll_interval} segundos...")
            try:
                procurar_found = False
                for iframe in driver.find_elements(By.TAG_NAME, "iframe"):
                    try:
                        driver.switch_to.frame(iframe)
                        WebDriverWait(driver, 2).until(
                            EC.element_to_be_clickable(
                                (By.XPATH, "//a[@href=\"javascript:doAcao('browse');\"]"))
                        ).click()
                        procurar_found = True
                        driver.switch_to.default_content()
                        break
                    except Exception:
                        driver.switch_to.default_content()

                if not procurar_found:
                    print(
                        "⚠️ Ícone 'Procurar' não foi encontrado para atualizar a lista.")

            except Exception as e:
                print(f"⚠️ Erro ao clicar em 'Procurar': {e}")
                driver.switch_to.default_content()

            time.sleep(poll_interval)

        if not download_found:
            raise Exception(
                "Tempo esgotado. O link de download não foi encontrado.")

        print("🕒 Aguardando o arquivo .zip ser baixado...")
        latest_zip = None
        zips_before = set(glob.glob(os.path.join(download_dir, '*.zip')))

        for _ in range(120):
            zips_after = set(glob.glob(os.path.join(download_dir, '*.zip')))
            new_zips = zips_after - zips_before
            if new_zips:
                latest_zip = max(new_zips, key=os.path.getctime)
                break
            time.sleep(0.5)

        if not latest_zip:
            raise FileNotFoundError(
                "❌ Download do arquivo .zip falhou ou demorou demais.")

        print(f"📦 Extraindo de: {latest_zip}")
        with zipfile.ZipFile(latest_zip, 'r') as zip_ref:
            xls_files = [f for f in zip_ref.namelist() if f.lower().endswith('.xls')]
            if not xls_files:
                raise FileNotFoundError(
                    "❌ Nenhum arquivo .xls encontrado no .zip.")
            xls_file = xls_files[0]
            zip_ref.extract(xls_file, download_dir)

        xls_path = os.path.join(download_dir, xls_file)

        for _ in range(10):
            if os.path.exists(xls_path):
                break
            time.sleep(0.5)
        if not os.path.exists(xls_path):
            raise FileNotFoundError(
                f"❌ Arquivo .xls extraído não foi encontrado: {xls_path}")

        print("🚀 Otimizando a planilha...")
        excel = None
        try:
            excel = win32.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False

            wb = excel.Workbooks.Open(xls_path)
            xlsx_filename = "Planilha de pendentes.xlsx"
            dest_path = os.path.join(desktop_dir, xlsx_filename)
            counter = 1
            while os.path.exists(dest_path):
                xlsx_filename = f"Planilha de pendentes_{counter}.xlsx"
                dest_path = os.path.join(desktop_dir, xlsx_filename)
                counter += 1
            wb.SaveAs(dest_path, FileFormat=51)
            wb.Close(False)
            wb = excel.Workbooks.Open(dest_path)
            ws = wb.Worksheets(1)
            ws.Activate()

            df_controle = app_instance.df
            whitelist_empresas = set(
                df_controle.iloc[:, 0].dropna().astype(str).str.upper())
            print(
                f"✅ Whitelist dinâmica criada com {len(whitelist_empresas)} empresas.")

            max_row = ws.UsedRange.Rows.Count
            for i in range(max_row, 4, -1):
                empresa_cell = ws.Cells(i, 1)
                empresa = str(empresa_cell.Value).strip().upper(
                ) if empresa_cell.Value else ""
                row_range = ws.Rows(i)

                if empresa not in whitelist_empresas:
                    row_range.Delete()

            used_range = ws.UsedRange
            used_range.Borders.LineStyle = 1
            used_range.Borders.Weight = 2
            used_range.HorizontalAlignment = -4108
            used_range.VerticalAlignment = -4108
            used_range.WrapText = True
            ws.Columns.AutoFit()
            ws.Rows.AutoFit()

            for r in range(1, ws.UsedRange.Rows.Count + 1):
                if ws.Rows(r).RowHeight < 23:
                    ws.Rows(r).RowHeight = 23

            ws.Range("A1").Select()
            wb.Save()

            excel.Visible = True
            excel.DisplayAlerts = True
            excel.ActiveWindow.Zoom = 100
            excel.ActiveWindow.ScrollRow = 1
            excel.ActiveWindow.ScrollColumn = 1

            start_time_excel = time.time()
            excel_maximized = False
            while time.time() - start_time_excel < 10:
                try:
                    app = Application(backend="uia").connect(
                        process=excel.ProcessID)
                    window = app.top_window()
                    if window.is_visible():
                        window.set_focus().set_foreground().maximize()
                        excel_maximized = True
                        print("✅ Janela do Excel maximizada com sucesso.")
                        break
                except Exception:
                    time.sleep(0.25)

            if not excel_maximized:
                print(
                    "⚠️ Aviso: Não foi possível focar e maximizar a janela do Excel após 10 segundos.")

            app_instance.root.after(0, app_instance.root.iconify)

        except Exception as e:
            if excel:
                excel.Quit()
            messagebox.showerror("Erro", f"Falha ao processar a planilha: {e}")
            raise e

        with open(TXT_ULTIMO_GERAR, "w", encoding="utf-8") as f:
            f.write(datetime.now().strftime("%d/%m/%Y"))
        hide_file(TXT_ULTIMO_GERAR)
        app_instance.root.after(0, btn_widget.pack_forget)

    except Exception as e:
        print(f"❌ Erro: {e}")
        messagebox.showerror("Erro", f"Falha ao gerar planilha: {e}")
    finally:
        print("🟢 Processo finalizado. Fechando navegador em segundo plano.")
        driver.quit()
        pythoncom.CoUninitialize()
        if status_banner:
            app_instance.root.after(0, status_banner.destroy)
        try:
            shutil.rmtree(temp_profile_dir)
            print(
                f"🧹 Perfil temporário '{temp_profile_dir}' limpo com sucesso.")
        except Exception as e:
            print(f"⚠️ Aviso: Falha ao limpar o perfil temporário: {e}")


class GerarPlanilhaWindow(ctk.CTkToplevel):
    def __init__(self, parent, usuario_atual, app_instance, btn_widget):
        super().__init__(parent)
        self.parent = parent
        self.usuario_atual = usuario_atual
        self.app_instance = app_instance
        self.btn_widget = btn_widget
        app_instance.gerar_planilha_window = self

        self.title("E.A.I. - GERAÇÃO DA PLANILHA DE PENDENTES")
        largura, altura = 600, 500
        x = (self.winfo_screenwidth() // 2) - (largura // 2)
        y = (self.winfo_screenheight() // 2) - (altura // 2)
        self.geometry(f"{largura}x{altura}+{x}+{y}")
        self.resizable(False, False)
        self.after(200, lambda: self.iconify()) # Workaround para aparecer na frente
        self.after(400, lambda: self.deiconify())
        self.focus()

        self.senha_visible = tk.BooleanVar(value=False)
        self.id_visible = tk.BooleanVar(value=False)

        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.pack(expand=True, pady=20, padx=20, fill="both")

        ctk.CTkLabel(frame, text="USUÁRIO SOC:", font=FONT_LABEL).grid(row=0, column=0, sticky="e", padx=10, pady=10)
        self.entry_usuario = ctk.CTkEntry(frame, font=FONT_ENTRY, width=250, height=35)
        self.entry_usuario.grid(row=0, column=1, sticky="w", padx=10, pady=10)

        ctk.CTkLabel(frame, text="SENHA SOC:", font=FONT_LABEL).grid(row=1, column=0, sticky="e", padx=10, pady=10)
        self.entry_senha = ctk.CTkEntry(frame, font=FONT_ENTRY, show='*', width=250, height=35)
        self.entry_senha.grid(row=1, column=1, sticky="w", padx=10, pady=10)

        self.senha_icon = ctk.CTkLabel(frame, text="🔒", font=("Segoe UI Emoji", 16), cursor="hand2")
        self.senha_icon.grid(row=1, column=2, sticky="w", padx=5)
        self.senha_icon.bind("<Button-1>", lambda e: self.toggle_visibility(self.entry_senha, self.senha_visible, self.senha_icon))

        ctk.CTkLabel(frame, text="ID SOC:", font=FONT_LABEL).grid(row=2, column=0, sticky="e", padx=10, pady=10)
        self.entry_id = ctk.CTkEntry(frame, font=FONT_ENTRY, show='*', width=250, height=35)
        self.entry_id.grid(row=2, column=1, sticky="w", padx=10, pady=10)

        self.id_icon = ctk.CTkLabel(frame, text="🔒", font=("Segoe UI Emoji", 16), cursor="hand2")
        self.id_icon.grid(row=2, column=2, sticky="w", padx=5)
        self.id_icon.bind("<Button-1>", lambda e: self.toggle_visibility(self.entry_id, self.id_visible, self.id_icon))

        self.caps_lock_label_gerar = ctk.CTkLabel(frame, text="", font=FONT_BOLD, text_color="red")
        self.caps_lock_label_gerar.grid(row=3, column=0, columnspan=3, sticky="we", padx=10, pady=(5, 0))
        self.app_instance.caps_lock_labels.append((self.caps_lock_label_gerar, "transparent"))

        self.btn_gerar = ctk.CTkButton(frame, text="GERAR PLANILHA DE PENDENTES", font=FONT_BUTTON, height=45, command=self.executar)
        self.btn_gerar.grid(row=4, column=0, columnspan=3, pady=30, padx=20, sticky="we")

        self.entry_usuario.bind("<Return>", lambda e: self.entry_senha.focus_set())
        self.entry_senha.bind("<Return>", lambda e: self.entry_id.focus_set())
        self.entry_id.bind("<Return>", lambda e: self.executar())
        self.entry_usuario.focus_set()

        rodape = ctk.CTkFrame(self, height=50, fg_color="transparent")
        rodape.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=10)

        voltar_btn = ctk.CTkButton(rodape, text="VOLTAR", font=FONT_BUTTON, width=100, fg_color="#E74C3C", hover_color="#C0392B", command=self.on_win_destroy)
        voltar_btn.pack(side=tk.LEFT)

        help_link = ctk.CTkLabel(rodape, text="ENCONTROU UM ERRO? CLIQUE AQUI", font=FONT_FORGOT_PASSWORD, text_color="#3498DB", cursor="hand2")
        help_link.pack(side=tk.RIGHT)
        help_link.bind("<Button-1>", abrir_link_ajuda)

    def toggle_visibility(self, entry_widget, var, label):
        is_currently_hidden = not var.get()
        if is_currently_hidden:
            entry_widget.configure(show="")
            label.configure(text="🔓")
        else:
            entry_widget.configure(show="*")
            label.configure(text="🔒")
        var.set(is_currently_hidden)

    def executar(self, event=None):
        soc_usuario = self.entry_usuario.get().strip()
        soc_senha = self.entry_senha.get().strip()
        soc_id = self.entry_id.get().strip()

        if not soc_usuario or not soc_senha or not soc_id:
            messagebox.showwarning("E.A.I. - ATENÇÃO", "POR FAVOR, PREENCHA TODOS OS CAMPOS (USUÁRIO, SENHA E ID).", parent=self)
            return
        if not soc_id.isdigit():
            messagebox.showwarning("E.A.I. - ATENÇÃO", "O ID DEVE CONTER APENAS NÚMEROS.", parent=self)
            return

        self.destroy()
        status_banner = self.app_instance.mostrar_status_geracao()
        threading.Thread(target=executar_gerar_planilha, args=(
            self.usuario_atual, soc_usuario, soc_senha, soc_id, self.app_instance, self.btn_widget, status_banner), daemon=True).start()

    def on_win_destroy(self):
        # Remove from caps lock labels list
        self.app_instance.caps_lock_labels = [item for item in self.app_instance.caps_lock_labels if item[0] != self.caps_lock_label_gerar]
        self.app_instance.gerar_planilha_window = None
        self.destroy()


def centralizar_janela(janela, largura=500, altura=500):
    janela.update_idletasks()
    largura_tela = janela.winfo_screenwidth()
    altura_tela = janela.winfo_screenheight()
    x = (largura_tela // 2) - (largura // 2)
    y = (altura_tela // 2) - (altura // 2)
    janela.geometry(f"{largura}x{altura}+{x}+{y}")

# ================== UTILITÁRIOS ==================

def centralizar(janela):
    janela.update_idletasks()
    largura = janela.winfo_width()
    altura = janela.winfo_height()
    x = (janela.winfo_screenwidth() // 2) - (largura // 2)
    y = (janela.winfo_screenheight() // 2) - (altura // 2)
    janela.geometry(f"{largura}x{altura}+{x}+{y}")

def normalize_key(s: str) -> str:
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return s.lower()

def safe_email_field(val) -> str:
    if val is None:
        return ""
    try:
        if isinstance(val, float) and pd.isna(val):
            return ""
    except Exception:
        pass
    s = str(val).strip()
    if s.lower() in {"nan", "none", "null", ""}:
        return ""
    return s

def limpar_caminho(p: str) -> str:
    if p is None:
        return ""
    s = str(p).strip().strip('"').strip("'")
    return s

def extrair_nome_empresa_do_arquivo(caminho: str) -> str:
    nome = os.path.basename(caminho)
    nome = re.sub(r"(?i)listagem de funcionarios\s*-\s*", "", nome)
    nome = re.sub(r"\.(xlsx|xlsm|xls)$", "", nome, flags=re.I)
    return nome.strip()

# ================== DADOS ==================

def carregar_controle(usuario):
    if usuario == "thailany":
        caminho_controle = CAMINHO_CONTROLE_THAILANY
    else:
        caminho_controle = CAMINHO_CONTROLE_VICTOR
    if not os.path.isfile(caminho_controle):
        raise FileNotFoundError(f"Planilha não encontrada: {caminho_controle}")
    df = pd.read_excel(caminho_controle, sheet_name=0, engine="openpyxl")
    print(f"Shape do DataFrame: {df.shape}")
    if df.shape[1] <= IDX_AK_LABEL:
        raise ValueError(
            f"A planilha tem {df.shape[1]} colunas, mas o código espera pelo menos {IDX_AK_LABEL + 1} colunas.")
    return df

def carregar_txt(nome_arquivo):
    registros = {}
    if os.path.isfile(nome_arquivo):
        with open(nome_arquivo, "r", encoding="utf-8") as f:
            for linha in f:
                linha = linha.strip()
                if "|" in linha:
                    empresa, data = linha.split("|", 1)
                    registros[empresa] = data
    return registros

def registrar_txt(nome_arquivo, empresa):
    hoje = datetime.now().strftime("%d/%m/%Y")
    if not os.path.isfile(nome_arquivo):
        with open(nome_arquivo, "w", encoding="utf-8") as f:
            f.write("")
        hide_file(nome_arquivo)
    with open(nome_arquivo, "a", encoding="utf-8") as f:
        f.write(f"{empresa}|{hoje}\n")
    hide_file(nome_arquivo)

# ================== INTRO PATINHO ==================

def mostrar_intro(root, callback):
    import sys
    if getattr(sys, 'frozen', False):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.dirname(__file__)

    caminho_patinho = None
    caminho_patinho_1 = os.path.join(base_path, "patinho.gif")
    if os.path.isfile(caminho_patinho_1):
        caminho_patinho = caminho_patinho_1
    elif os.path.isfile(CAMINHO_PATINHO_2):
        caminho_patinho = CAMINHO_PATINHO_2

    intro = tk.Toplevel(root)
    intro.attributes("-fullscreen", True)
    intro.configure(bg="white")

    if caminho_patinho:
        try:
            img = tk.PhotoImage(file=caminho_patinho)
            lbl_img = tk.Label(intro, image=img, bg="white")
            lbl_img.image = img
            lbl_img.pack(expand=True)
        except Exception as e:
            lbl_txt = tk.Label(
                intro,
                text=f"[Erro carregando imagem: {e}]",
                fg="black",
                bg="white",
                font=FONT_TITLE)
            lbl_txt.pack(expand=True)
    else:
        lbl_txt = tk.Label(
            intro,
            text="[Imagem patinho.gif não encontrada]",
            fg="black",
            bg="white",
            font=FONT_TITLE)
        lbl_txt.pack(expand=True)
    intro.after(2000, lambda: (intro.destroy(), callback()))

# ================== ENVIO ==================

def enviar_email_label(label, df, usuario):
    try:
        pythoncom.CoInitialize()
        outlook = win32.Dispatch("Outlook.Application")
        session = outlook.Session
        row = get_row_by_label(df, label)
        if row is None:
            raise Exception(
                f"Botão/label '{label}' não encontrado na planilha.")
        email_to = safe_email_field(row.iloc[IDX_EMAIL_TO])
        email_cc = safe_email_field(row.iloc[IDX_EMAIL_CC])
        if not email_to:
            raise Exception(
                f"E-mail de destino não configurado para '{label}'.")
        anexos = get_anexos_from_row(row, usuario)
        if not anexos:
            raise Exception(f"Anexos não encontrados para '{label}'.")
        assunto = get_assunto_from_row(row)
        if not assunto:
            nomes_empresas = [
                extrair_nome_empresa_do_arquivo(a) for a in anexos]
            assunto = "Listagem de Funcionários - " + \
                " / ".join(nomes_empresas)
        mail = outlook.CreateItem(0)
        try:
            accounts = session.Accounts
            email_conta = EMAIL_CONTA_PADRAO_THAILANY if usuario == "thailany" else EMAIL_CONTA_PADRAO_VICTOR
            account_found = False
            for acc in accounts:
                if str(acc.SmtpAddress).lower() == email_conta.lower():
                    mail._oleobj_.Invoke(64209, 0, 8, 0, acc)
                    account_found = True
                    break
            if not account_found:
                raise Exception(
                    f"Conta de e-mail {email_conta} não encontrada no Outlook.")
        except Exception as e:
            raise Exception(f"Erro ao selecionar conta de envio: {e}")
        mail.To = email_to
        if email_cc:
            mail.CC = email_cc
        mail.Subject = assunto
        mail.Display(False)
        time.sleep(0.25)
        assinatura_html = mail.HTMLBody or ""
        corpo_html = (
            "<html><body>"
            "<div style='font-family: Calibri; font-size: 11pt;'>"
            "<p>Bom dia,</p>"
            "<p>Segue listagem de funcionários que estão Ativos em nosso sistema. "
            "Por gentileza confirmar se a listagem está correta, conferindo Nome, Cargo e Situação (Ativo/Afastado/Pendente).</p>"
            "</div>"
            "</body></html>"
        )
        mail.HTMLBody = corpo_html + assinatura_html
        for anexo in anexos:
            mail.Attachments.Add(os.path.abspath(anexo))
        mail.Send()
        registrar_txt(TXT_ENVIO, label)
        return True, None
    except Exception as e:
        return False, str(e)
    finally:
        pythoncom.CoUninitialize()

# ================== COBRANÇA ==================

def cobrar_email_label(label, df, usuario):
    try:
        pythoncom.CoInitialize()
        outlook = win32.Dispatch("Outlook.Application")
        ns = outlook.GetNamespace("MAPI")
        email_conta = EMAIL_CONTA_PADRAO_THAILANY if usuario == "thailany" else EMAIL_CONTA_PADRAO_VICTOR
        account_found = False
        for account in ns.Accounts:
            if str(account.SmtpAddress).lower() == email_conta.lower():
                sent_items = account.DeliveryStore.GetDefaultFolder(
                    5)
                account_found = True
                print(f"✅ Conta encontrada: {email_conta}")
                break
        if not account_found:
            raise Exception(
                f"Conta de e-mail {email_conta} não encontrada no Outlook.")
        items = sent_items.Items
        items.Sort("[SentOn]", True)
        row = get_row_by_label(df, label)
        if row is None:
            raise Exception(
                f"Botão/label '{label}' não encontrado na planilha.")
        email_to = safe_email_field(row.iloc[IDX_EMAIL_TO])
        email_cc = safe_email_field(row.iloc[IDX_EMAIL_CC])
        if not email_to:
            raise Exception(
                f"E-mail de destino não configurado para '{label}'.")
        anexos = get_anexos_from_row(row, usuario)
        if not anexos:
            raise Exception(f"Anexos não encontrados para '{label}'.")
        assunto = get_assunto_from_row(row)
        if not assunto:
            nomes_empresas = [extrair_nome_empresa_do_arquivo(a)
                              for a in anexos]
            assunto = "Listagem de Funcionários - " + " / ".join(nomes_empresas)
        assunto_normalizado = unicodedata.normalize(
            "NFD", assunto).encode('ascii', 'ignore').decode('ascii').lower()
        assunto_normalizado = re.sub(
            r'^(re|fwd|fw):\s*', '', assunto_normalizado, flags=re.I)
        assunto_normalizado = re.sub(
            r'\s+', ' ', assunto_normalizado.strip())
        print(
            f"🔍 Procurando e-mail com assunto normalizado: '{assunto_normalizado}'")
        hoje = datetime.now()
        last_mail = None
        for itm in items:
            try:
                sent_on = itm.SentOn
                if sent_on.month != hoje.month or sent_on.year != hoje.year:
                    continue
                subj = (itm.Subject or "").lower()
                subj = unicodedata.normalize(
                    "NFD", subj).encode('ascii', 'ignore').decode('ascii')
                subj = re.sub(r'^(re|fwd|fw):\s*', '', subj, flags=re.I)
                subj = re.sub(r'\s+', ' ', subj.strip())
                sender = (itm.SenderEmailAddress or "").lower()
                if email_conta.lower() in sender and assunto_normalizado in subj:
                    last_mail = itm
                    break
            except Exception:
                continue
        if last_mail is None:
            raise Exception(
                f"Não foi possível localizar e-mail anterior do mês corrente com assunto contendo: '{assunto}'.")
        reply = last_mail.ReplyAll()
        reply.Subject = last_mail.Subject
        reply.To = email_to
        if email_cc:
            reply.CC = email_cc
        corpo_cobranca = (
            "<html><body>"
            "<div style='font-family: Calibri; font-size: 11pt;'>"
            "<p>Prezados,</p><p>Algum retorno?</p>"
            "</div>"
            "</body></html>"
        )
        reply.HTMLBody = corpo_cobranca + (reply.HTMLBody or "")
        for anexo in anexos:
            reply.Attachments.Add(os.path.abspath(anexo))
        try:
            accounts = ns.Accounts
            account_found = False
            for acc in accounts:
                if str(acc.SmtpAddress).lower() == email_conta.lower():
                    reply._oleobj_.Invoke(64209, 0, 8, 0, acc)
                    account_found = True
                    break
            if not account_found:
                raise Exception(
                    f"Conta de e-mail {email_conta} não encontrada para envio.")
        except Exception as e:
            raise Exception(
                f"Erro ao selecionar conta de envio para cobrança: {e}")
        reply.Send()
        registrar_txt(TXT_COBRANCA, label)
        return True, None
    except Exception as e:
        return False, str(e)
    finally:
        pythoncom.CoUninitialize()

# ================== OCULTAR/EXIBIR ARQUIVOS TXT ==================

def hide_file(file_path):
    try:
        if os.path.exists(file_path):
            win32api.SetFileAttributes(
                file_path, win32con.FILE_ATTRIBUTE_HIDDEN)
    except Exception:
        pass

def unhide_file(file_path):
    try:
        if os.path.exists(file_path):
            win32api.SetFileAttributes(
                file_path, win32con.FILE_ATTRIBUTE_NORMAL)
    except Exception:
        pass

# ================== APLICAÇÃO (UI) ==================

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("E.A.I. - E-SOCIAL ARTIFICIAL INTELLIGENCE")
        self.root.geometry("700x550")
        centralizar(self.root)

        for txt_file in [
                TXT_ULTIMO_ENVIO,
                TXT_ENVIO,
                TXT_COBRANCA,
                TXT_PROCURACOES,
                TXT_ULTIMO_AJUSTE,
                TXT_ULTIMO_GERAR]:
            unhide_file(txt_file)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

        self.usuario_atual = None
        self.df = pd.DataFrame()
        self.labels = []
        self.enviados = 0
        self.cobrados = 0
        self.baixados = 0
        self.ajustadas = 0
        self.procuracoes_enviadas = 0
        self.frame_empresas_envio = None
        self.frame_empresas_cobranca = None
        self.frame_empresas_procuracoes = None
        self.password_visible = False
        self.show_envio_cobranca_button = True
        self.caps_lock_labels = []
        self.gerar_planilha_window = None

        self._build_login_screen()
        self.check_caps_lock()

    def check_caps_lock(self):
        try:
            is_caps_on = win32api.GetKeyState(win32con.VK_CAPITAL) & 1
            for label, _ in self.caps_lock_labels:
                try:
                    if is_caps_on:
                        label.configure(text="CAPS LOCK ATIVADA")
                    else:
                        label.configure(text="")
                except Exception:
                    pass
        except Exception:
            pass
        self.root.after(500, self.check_caps_lock)

    def _build_login_screen(self):
        self.overlay_login = ctk.CTkFrame(self.root, fg_color="transparent")
        self.overlay_login.place(relx=0, rely=0, relwidth=1, relheight=1)

        login_container = ctk.CTkFrame(self.overlay_login, width=450, height=450, corner_radius=15)
        login_container.pack_propagate(False)
        login_container.place(relx=0.5, rely=0.5, anchor="center")

        ctk.CTkLabel(login_container, text="E.A.I.", font=ctk.CTkFont(family=FONT_FAMILY, size=32, weight="bold"), text_color="#2C3E50").pack(pady=(30, 5))
        ctk.CTkLabel(login_container, text="Acesso ao Sistema", font=FONT_LABEL, text_color="#7F8C8D").pack(pady=(0, 25))

        self.usuario_entry = ctk.CTkEntry(login_container, placeholder_text="Usuário", font=FONT_ENTRY, width=320, height=45)
        self.usuario_entry.pack(pady=10)

        # Senha container para o ícone de olho
        senha_frame = ctk.CTkFrame(login_container, fg_color="transparent")
        senha_frame.pack(pady=10)

        self.senha_entry = ctk.CTkEntry(senha_frame, placeholder_text="Senha", show="*", font=FONT_ENTRY, width=320, height=45)
        self.senha_entry.pack(side="left")

        self.eye_icon_label = ctk.CTkLabel(senha_frame, text="🔒", font=("Segoe UI Emoji", 18), cursor="hand2")
        self.eye_icon_label.place(relx=0.92, rely=0.5, anchor="center")
        self.eye_icon_label.bind("<Button-1>", self.toggle_password_visibility)

        self.caps_lock_msg = ctk.CTkLabel(login_container, text="", font=(FONT_FAMILY, 11), text_color="#E74C3C")
        self.caps_lock_msg.pack(pady=2)
        self.caps_lock_labels.append((self.caps_lock_msg, "transparent"))

        forgot_btn = ctk.CTkLabel(login_container, text="Esqueci a minha senha", font=FONT_FORGOT_PASSWORD, text_color="#3498DB", cursor="hand2")
        forgot_btn.pack(pady=5)
        forgot_btn.bind("<Button-1>", self.abrir_janela_palavra_seguranca)

        btn_entrar = ctk.CTkButton(login_container, text="ENTRAR", font=FONT_BUTTON, width=320, height=50, corner_radius=8, command=self.validar_login)
        btn_entrar.pack(pady=25)

        footer_frame = ctk.CTkFrame(self.overlay_login, height=40, fg_color="transparent")
        footer_frame.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=10)

        ctk.CTkButton(footer_frame, text="SAIR", font=FONT_BUTTON, width=80, fg_color="#E74C3C", hover_color="#C0392B", command=self.root.destroy).pack(side=tk.LEFT)

        help_link = ctk.CTkLabel(footer_frame, text="ENCONTROU UM ERRO? CLIQUE AQUI", font=FONT_FORGOT_PASSWORD, text_color="#3498DB", cursor="hand2")
        help_link.pack(side=tk.RIGHT)
        help_link.bind("<Button-1>", abrir_link_ajuda)

        self.usuario_entry.bind("<Return>", lambda e: self.senha_entry.focus_set())
        self.senha_entry.bind("<Return>", lambda e: self.validar_login(e))
        self.usuario_entry.focus_set()

    def tela_login(self):
        for widget in self.root.winfo_children():
            widget.destroy()
        self.root.title("E.A.I. - E-SOCIAL ARTIFICIAL INTELLIGENCE")
        self.root.geometry("700x550")
        centralizar(self.root)
        self.usuario_atual = None
        self.df = pd.DataFrame()
        self.labels = []
        self.show_envio_cobranca_button = True
        self.caps_lock_labels = []
        self._build_login_screen()

    def toggle_password_visibility(self, event=None):
        if self.password_visible:
            self.senha_entry.configure(show="*")
            self.eye_icon_label.configure(text="🔒")
            self.password_visible = False
        else:
            self.senha_entry.configure(show="")
            self.eye_icon_label.configure(text="🔓")
            self.password_visible = True

    def on_close(self):
        for txt_file in [
                TXT_ULTIMO_ENVIO,
                TXT_ENVIO,
                TXT_COBRANCA,
                TXT_PROCURACOES,
                TXT_ULTIMO_AJUSTE,
                TXT_ULTIMO_GERAR]:
            hide_file(txt_file)
        self.root.destroy()

    def abrir_janela_palavra_seguranca(self, event=None):
        initial_usuario = self.usuario_entry.get().strip().lower()
        win = ctk.CTkToplevel(self.root)
        win.title("E.A.I. - RECUPERAÇÃO DE SENHA")
        largura, altura = 600, 500
        x = (self.winfo_screenwidth() // 2) - (largura // 2)
        y = (self.winfo_screenheight() // 2) - (altura // 2)
        win.geometry(f"{largura}x{altura}+{x}+{y}")
        win.resizable(False, False)
        win.after(200, lambda: win.focus())

        tentativas_restantes = 3

        container = ctk.CTkFrame(win, corner_radius=15)
        container.pack(expand=True, padx=40, pady=40, fill="both")

        ctk.CTkLabel(container, text="Recuperação de Senha", font=FONT_TITLE).pack(pady=20)

        ctk.CTkLabel(container, text="USUÁRIO:", font=FONT_LABEL).pack(anchor="w", padx=50)
        usuario_rec_entry = ctk.CTkEntry(container, font=FONT_ENTRY, width=300)
        usuario_rec_entry.pack(pady=(0, 15))

        ctk.CTkLabel(container, text="PALAVRA DE SEGURANÇA:", font=FONT_LABEL).pack(anchor="w", padx=50)

        palavra_frame = ctk.CTkFrame(container, fg_color="transparent")
        palavra_frame.pack()
        palavra_entry = ctk.CTkEntry(palavra_frame, font=FONT_ENTRY, show="*", width=300)
        palavra_entry.pack(side="left")

        palavra_visible = tk.BooleanVar(value=False)
        palavra_icon = ctk.CTkLabel(palavra_frame, text="🔒", font=("Segoe UI Emoji", 16), cursor="hand2")
        palavra_icon.place(relx=0.92, rely=0.5, anchor="center")

        def toggle_palavra_visibility():
            is_hidden = not palavra_visible.get()
            palavra_entry.configure(show="" if is_hidden else "*")
            palavra_icon.configure(text="🔓" if is_hidden else "🔒")
            palavra_visible.set(is_hidden)

        palavra_icon.bind("<Button-1>", lambda e: toggle_palavra_visibility())

        if initial_usuario:
            usuario_rec_entry.insert(0, initial_usuario)
            usuario_rec_entry.configure(state="readonly")
            palavra_entry.focus_set()
        else:
            usuario_rec_entry.focus_set()

        contador_label = ctk.CTkLabel(container, text=f"RESTAM {tentativas_restantes} TENTATIVAS.", font=FONT_BOLD, text_color="#E74C3C")
        contador_label.pack(pady=10)

        def verificar():
            nonlocal tentativas_restantes
            usuario = usuario_rec_entry.get().strip().lower()
            palavra_digitada = palavra_entry.get().strip().lower()

            if not usuario:
                messagebox.showwarning("ATENÇÃO", "POR FAVOR, INFORME O USUÁRIO.", parent=win)
                return

            if not palavra_digitada:
                self._show_palavra_seguranca_dialog(win, usuario)
                return

            PALAVRAS_SECRETAS = {"victor": "creatina", "thailany": "spoc"}

            if usuario in PALAVRAS_SECRETAS and palavra_digitada == PALAVRAS_SECRETAS[usuario]:
                messagebox.showinfo("SENHA RECUPERADA", f"OLÁ, {usuario.upper()}, SUA SENHA É: 1001", parent=win)
                win.destroy()
            else:
                tentativas_restantes -= 1
                if tentativas_restantes > 0:
                    contador_label.configure(text=f"RESTAM {tentativas_restantes} TENTATIVAS.")
                    messagebox.showwarning("DADOS INCORRETOS", "USUÁRIO OU PALAVRA DE SEGURANÇA INCORRETA.", parent=win)
                else:
                    messagebox.showerror("TENTATIVAS ESGOTADAS", "VOCÊ ESGOTOU SUAS TENTATIVAS.", parent=win)
                    webbrowser.open("https://wa.me/551221239207?text=N%C3%A3o%20lembro%20a%20minha%20senha%2Fpalavra%20de%20seguran%C3%A7a%20do%20sistema%20E.A.I.%20Pode%20me%20ajudar%3F")
                    win.destroy()

        ctk.CTkButton(container, text="RECUPERAR SENHA", font=FONT_BUTTON, width=300, height=45, command=verificar).pack(pady=20)

        rodape = ctk.CTkFrame(win, height=40, fg_color="transparent")
        rodape.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=10)
        ctk.CTkButton(rodape, text="VOLTAR", font=FONT_BUTTON, width=100, fg_color="#E74C3C", command=win.destroy).pack(side=tk.LEFT)

    def _enviar_uma_procuracao_thread(self, empresa, btn_widget):
        success, err = enviar_email_procuracao(empresa, self.df, self.usuario_atual)
        def ui_update():
            if success:
                try: btn_widget.destroy()
                except Exception: pass
                self.procuracoes_enviadas += 1
            else:
                messagebox.showerror("Erro de Envio", f"Falha ao enviar e-mail para '{empresa}': {err}")
            if self.frame_empresas_procuracoes and self.frame_empresas_procuracoes.winfo_exists():
                btns = [w for w in self.frame_empresas_procuracoes.winfo_children() if isinstance(w, ctk.CTkButton)]
                self.contador_label_procuracoes.configure(text=f"QUANTIDADE DE PROCURAÇÕES VENCIDAS / A VENCER: {len(btns)}")
        self.root.after(0, ui_update)

    def _enviar_todas_procuracoes_thread(self):
        if not self.frame_empresas_procuracoes: return
        widgets = [w for w in self.frame_empresas_procuracoes.winfo_children() if isinstance(w, ctk.CTkButton)]
        for widget in widgets:
            empresa = widget.cget("text")
            self._enviar_uma_procuracao_thread(empresa, widget)
            time.sleep(0.5)

    def tela_procuracoes(self):
        win = ctk.CTkToplevel(self.root)
        win.title("E.A.I. - PROCURAÇÕES")
        win.after(200, lambda: win.state('zoomed'))

        frame_top = ctk.CTkFrame(win, fg_color="transparent")
        frame_top.pack(side=tk.TOP, fill=tk.X, pady=20)
        ctk.CTkLabel(frame_top, text="PROCURAÇÕES VENCIDAS", font=FONT_TITLE).pack()

        ctk.CTkButton(frame_top, text="ENVIAR TODOS", font=FONT_BUTTON, fg_color="#1F5FBF", hover_color="#154360", command=lambda: threading.Thread(target=self._enviar_todas_procuracoes_thread, daemon=True).start()).pack(pady=10)

        hoje = datetime.now()
        registros_procuracoes = carregar_txt(TXT_PROCURACOES)
        botoes_a_criar = []

        for index, row in self.df.iloc[2:].iterrows():
            nome_empresa = str(row.iloc[0])
            validade_cell = row.iloc[IDX_AC_VALIDADE]
            status_cell = str(row.iloc[IDX_F_BLOQUEADA])
            if pd.isna(nome_empresa) or not nome_empresa.strip() or "BLOQUEADA" in status_cell.upper(): continue
            if nome_empresa in registros_procuracoes and registros_procuracoes[nome_empresa] == hoje.strftime("%d/%m/%Y"): continue

            if isinstance(validade_cell, datetime): validade_data = validade_cell
            elif isinstance(validade_cell, str):
                try: validade_data = datetime.strptime(validade_cell, "%d/%m/%Y")
                except ValueError: continue
            else: continue
            if validade_data < hoje: botoes_a_criar.append(nome_empresa)

        botoes_a_criar.sort(key=lambda x: normalize_key(x))
        self.contador_label_procuracoes = ctk.CTkLabel(frame_top, text=f"QUANTIDADE DE PROCURAÇÕES VENCIDAS / A VENCER: {len(botoes_a_criar)}", font=FONT_LABEL)
        self.contador_label_procuracoes.pack()

        scroll_container = ctk.CTkScrollableFrame(win, orientation="horizontal", label_text="Empresas com Pendência")
        scroll_container.pack(fill="both", expand=True, padx=20, pady=10)
        self.frame_empresas_procuracoes = scroll_container

        # Custom grid inside scrollable frame
        inner_frame = ctk.CTkFrame(scroll_container, fg_color="transparent")
        inner_frame.pack()
        self.frame_empresas_procuracoes = inner_frame

        for i, nome_empresa in enumerate(botoes_a_criar):
            r = i % MAX_LINHAS_COLUNA
            c = i // MAX_LINHAS_COLUNA
            btn = ctk.CTkButton(inner_frame, text=nome_empresa, font=FONT_BUTTON, width=300, height=80, corner_radius=10)
            btn.configure(command=lambda emp=nome_empresa, b=btn: threading.Thread(target=self._enviar_uma_procuracao_thread, args=(emp, b), daemon=True).start())
            btn.grid(row=r, column=c, padx=10, pady=10)

        rodape = ctk.CTkFrame(win, height=50, fg_color="transparent")
        rodape.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=10)
        ctk.CTkButton(rodape, text="VOLTAR", font=FONT_BUTTON, width=100, fg_color="#E74C3C", command=win.destroy).pack(side=tk.LEFT)

    def _has_envio_actions(self):
        registros_envio = carregar_txt(TXT_ENVIO)
        return any(not self._enviado_no_mes(lbl, registros_envio) and not self._has_ok_status(lbl) for lbl in self.labels)

    def _has_cobranca_actions(self):
        registros_cobranca = carregar_txt(TXT_COBRANCA)
        return any(not self._cobrado_hoje(lbl, registros_cobranca) and lbl not in EXCLUIR_NA_COBRANCA and not self._has_ok_status(lbl) for lbl in self.labels)

    def validar_login(self, event=None):
        usuario = self.usuario_entry.get().strip().lower()
        senha = self.senha_entry.get().strip()
        if not usuario or not senha:
            messagebox.showwarning("E.A.I. - ATENÇÃO", "POR FAVOR, PREENCHA OS DADOS.")
            return

        USUARIOS = {"victor": "1001", "thailany": "1001"}
        if usuario in USUARIOS and senha == USUARIOS[usuario]:
            self.usuario_atual = usuario
            try:
                self.df = carregar_controle(usuario)
                self.labels = listar_labels(self.df)
            except Exception as e:
                messagebox.showerror("E.A.I. - ERRO", f"FALHA AO CARREGAR PLANILHA: {e}")
                return
            self.overlay_login.destroy()
            messagebox.showinfo("E.A.I. - BEM-VINDO", f"LOGIN REALIZADO COM SUCESSO, {usuario.upper()}!")
            self.tela_inicial()
        else:
            self._show_login_error_dialog()

    def _show_login_error_dialog(self):
        dialog = ctk.CTkToplevel(self.root)
        dialog.title("E.A.I. - DADOS INCORRETOS")
        largura, altura = 500, 300
        x = (self.winfo_screenwidth() // 2) - (largura // 2)
        y = (self.winfo_screenheight() // 2) - (altura // 2)
        dialog.geometry(f"{largura}x{altura}+{x}+{y}")
        dialog.resizable(False, False)
        dialog.after(200, lambda: dialog.focus())

        ctk.CTkLabel(dialog, text="USUÁRIO E/OU SENHA INCORRETOS.\nGOSTARIA DE TENTAR NOVAMENTE OU REDEFINIR A SENHA?", font=FONT_LABEL, justify="center").pack(pady=40)

        btn_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        btn_frame.pack()

        ctk.CTkButton(btn_frame, text="TENTAR NOVAMENTE", font=FONT_BUTTON, command=dialog.destroy).pack(side="left", padx=10)
        ctk.CTkButton(btn_frame, text="REDEFINIR SENHA", font=FONT_BUTTON, fg_color="#F39C12", hover_color="#D68910", command=lambda: (dialog.destroy(), self.abrir_janela_palavra_seguranca())).pack(side="left", padx=10)

    def _show_palavra_seguranca_dialog(self, parent_win, usuario):
        dialog = ctk.CTkToplevel(parent_win)
        dialog.title("E.A.I. - ATENÇÃO")
        dialog.geometry("500x350")
        centralizar_janela(dialog, 500, 350)

        ctk.CTkLabel(dialog, text="PALAVRA DE SEGURANÇA NÃO DIGITADA", font=FONT_BOLD, text_color="white", fg_color="#E74C3C", corner_radius=5, height=40).pack(pady=30, padx=20, fill="x")

        ctk.CTkButton(dialog, text="DIGITAR NOVAMENTE", font=FONT_BUTTON, width=300, command=dialog.destroy).pack(pady=10)
        ctk.CTkButton(dialog, text="ESQUECI A PALAVRA", font=FONT_BUTTON, width=300, fg_color="#7F8C8D", command=lambda: (dialog.destroy(), webbrowser.open("https://wa.me/551221239207?text=N%C3%A3o%20lembro%20a%20minha%20senha%2Fpalavra%20de%20seguran%C3%A7a%20do%20sistema%20E.A.I.%20Pode%20me%20ajudar%3F"))).pack(pady=10)

    def mostrar_status_geracao(self):
        for widget in self.status_frame.winfo_children(): widget.destroy()
        status_banner = ctk.CTkFrame(self.status_frame, fg_color="#D6EAF8", corner_radius=10)
        status_banner.pack(fill="x", pady=10, padx=20)
        status_label = ctk.CTkLabel(status_banner, text="PROCESSO DE GERAÇÃO INICIADO...", font=FONT_BOLD, text_color="#1B4F72")
        status_label.pack(pady=15)
        self.root.after(10000, lambda: self._iniciar_countdown(status_banner, status_label, 180))
        return status_banner

    def _iniciar_countdown(self, banner, label, segundos):
        if segundos >= 0:
            label.configure(text=f"AGUARDE {segundos} SEGUNDOS, POR FAVOR")
            self.root.after(1000, lambda: self._iniciar_countdown(banner, label, segundos - 1))
        else:
            try: banner.destroy()
            except Exception: pass

    def salvar_planilhas_outlook(self, botao, usuario):
        try:
            botao.pack_forget()
            counter_frame = ctk.CTkFrame(self.status_frame, fg_color="#D5F5E3", corner_radius=10)
            counter_frame.pack(fill="x", pady=10, padx=20)
            counter_label = ctk.CTkLabel(counter_frame, text="ANALISANDO E-MAILS...", font=FONT_BOLD, text_color="#145A32")
            counter_label.pack(pady=15)

            if usuario == "thailany":
                email_conta = EMAIL_PLANILHAS_THAILANY
                pasta_destino = PASTA_PLANILHAS_THAILANY
            else:
                email_conta = EMAIL_CONTA_PADRAO_VICTOR
                pasta_destino = PASTA_PLANILHAS_VICTOR

            pythoncom.CoInitialize()
            outlook = win32.Dispatch("Outlook.Application")
            ns = outlook.GetNamespace("MAPI")
            account_found = False
            for account in ns.Accounts:
                if str(account.SmtpAddress).lower() == email_conta.lower():
                    inbox = account.DeliveryStore.GetDefaultFolder(6)
                    account_found = True
                    break
            if not account_found: raise Exception(f"Conta {email_conta} não encontrada.")

            items = inbox.Items
            items.Sort("[ReceivedTime]", True)
            os.makedirs(pasta_destino, exist_ok=True)

            total_planilhas = 0
            relevant_emails = []
            temp_dir = os.path.join(pasta_destino, "temp_zip_scan")
            os.makedirs(temp_dir, exist_ok=True)

            for itm in items:
                try:
                    assunto_norm = normalize_key(itm.Subject)
                    if "listagem de funcionarios" in assunto_norm and itm.Attachments.Count > 0:
                        relevant_emails.append(itm)
                except Exception: continue

            for itm in relevant_emails:
                for att in itm.Attachments:
                    if att.FileName.lower().endswith(".zip"):
                        temp_zip = os.path.join(temp_dir, att.FileName)
                        try:
                            att.SaveAsFile(temp_zip)
                            with zipfile.ZipFile(temp_zip, 'r') as zip_ref:
                                total_planilhas += len([f for f in zip_ref.namelist() if f.lower().endswith(".xlsx")])
                        finally:
                            if os.path.exists(temp_zip): os.remove(temp_zip)
            shutil.rmtree(temp_dir)

            baixados = 0
            for itm in relevant_emails:
                try:
                    nome_base = re.sub(r'[<>:"/\\|?*]', '', itm.Subject).strip()
                    for att in itm.Attachments:
                        if att.FileName.lower().endswith(".zip"):
                            caminho_zip_temp = os.path.join(pasta_destino, att.FileName)
                            att.SaveAsFile(caminho_zip_temp)
                            try:
                                with zipfile.ZipFile(caminho_zip_temp, 'r') as zip_ref:
                                    for zip_info in [f for f in zip_ref.namelist() if f.lower().endswith(".xlsx")]:
                                        nome_arquivo = f"{nome_base}.xlsx"
                                        caminho_final = os.path.join(pasta_destino, nome_arquivo)
                                        c = 1
                                        while os.path.exists(caminho_final):
                                            caminho_final = os.path.join(pasta_destino, f"{nome_base}_{c}.xlsx")
                                            c += 1
                                        with zip_ref.open(zip_info) as src, open(caminho_final, 'wb') as tgt:
                                            shutil.copyfileobj(src, tgt)
                                        baixados += 1
                                        self.root.after(0, lambda b=baixados, t=total_planilhas: counter_label.configure(text=f"PLANILHAS SALVAS: {b} DE {t}"))
                            finally: os.remove(caminho_zip_temp)
                except Exception: continue

            with open(TXT_ULTIMO_ENVIO, "w", encoding="utf-8") as f: f.write(datetime.now().strftime("%d/%m/%Y"))
            hide_file(TXT_ULTIMO_ENVIO)
            counter_frame.destroy()
            messagebox.showinfo("Sucesso", "Todas as planilhas foram salvas.")
        except Exception as e:
            if 'counter_frame' in locals(): counter_frame.destroy()
            messagebox.showerror("Erro", str(e))
        finally: pythoncom.CoUninitialize()

    def _ajustar_planilhas_thread(self, botao):
        botao.pack_forget()
        SENHA_PROTECAO = "spoc"
        pasta = PASTA_PLANILHAS_THAILANY if self.usuario_atual == "thailany" else PASTA_PLANILHAS_VICTOR
        counter_frame = ctk.CTkFrame(self.status_frame, fg_color="#FAD7A0", corner_radius=10)
        counter_frame.pack(fill="x", pady=10, padx=20)
        counter_label = ctk.CTkLabel(counter_frame, text="IDENTIFICANDO PLANILHAS...", font=FONT_BOLD, text_color="#7E5109")
        counter_label.pack(pady=15)

        try:
            all_files = [f for f in os.listdir(pasta) if f.endswith(".xlsx")]
            if not all_files:
                counter_frame.destroy()
                messagebox.showinfo("Aviso", "Nenhuma planilha encontrada.")
                return

            arquivos_para_ajustar = []
            pythoncom.CoInitialize()
            ex = win32.Dispatch("Excel.Application")
            ex.Visible, ex.DisplayAlerts = False, False
            try:
                for f in all_files:
                    try:
                        wb = ex.Workbooks.Open(os.path.join(pasta, f), ReadOnly=True, Password="")
                        wb.Close(); arquivos_para_ajustar.append(f)
                    except: continue
            finally: ex.Quit(); pythoncom.CoUninitialize()

            total = len(arquivos_para_ajustar)
            if total == 0:
                counter_frame.destroy()
                messagebox.showinfo("Aviso", "Planilhas já estão ajustadas.")
                return

            ajustadas = 0
            hoje = datetime.now()
            for arquivo in arquivos_para_ajustar:
                caminho = os.path.join(pasta, arquivo)
                wb = openpyxl.load_workbook(caminho)
                for ws in wb.worksheets:
                    ws.delete_rows(2, 1)
                    merged = list(ws.merged_cells.ranges)
                    to_del = []
                    for r in range(2, ws.max_row + 1):
                        is_m = any(r >= m.min_row and r <= m.max_row for m in merged)
                        if not is_m:
                            status = ws.cell(r, 6).value
                            data_v = ws.cell(r, 7).value
                            if status and str(status).strip().lower() == "pendente":
                                try:
                                    if not isinstance(data_v, datetime): data_v = datetime.strptime(str(data_v), "%d/%m/%Y")
                                    if data_v.month == hoje.month and data_v.year == hoje.year: to_del.append(r)
                                except: continue
                    for r in sorted(to_del, reverse=True): ws.delete_rows(r, 1)
                    thin = Side(border_style="thin", color="000000")
                    border = Border(top=thin, left=thin, right=thin, bottom=thin)
                    for r in range(1, ws.max_row + 1):
                        for c in range(1, ws.max_column + 1):
                            if not any(r >= m.min_row and r <= m.max_row and c >= m.min_col and c <= m.max_col for m in ws.merged_cells.ranges):
                                cell = ws.cell(r, c)
                                cell.border = border
                                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                cell.number_format = 'DD/MM/YYYY' if c == 7 else '@'
                wb.save(caminho)
                ajustadas += 1
                self.root.after(0, lambda a=ajustadas: counter_label.configure(text=f"AJUSTADAS: {a} DE {total}"))

            pythoncom.CoInitialize()
            ex = win32.Dispatch("Excel.Application")
            ex.Visible, ex.DisplayAlerts = False, False
            for f in arquivos_para_ajustar:
                try:
                    wb = ex.Workbooks.Open(os.path.join(pasta, f))
                    for ws in wb.Worksheets:
                        ws.Columns.AutoFit(); ws.Rows.AutoFit()
                        ws.Protect(SENHA_PROTECAO)
                    wb.Protect(SENHA_PROTECAO); wb.Save(); wb.Close()
                except: continue
            ex.Quit(); pythoncom.CoUninitialize()

            with open(TXT_ULTIMO_AJUSTE, "w", encoding="utf-8") as f: f.write(datetime.now().strftime("%d/%m/%Y"))
            hide_file(TXT_ULTIMO_AJUSTE)
            counter_frame.destroy()
            messagebox.showinfo("Sucesso", f"Planilhas ajustadas. Senha: {SENHA_PROTECAO}")
        except Exception as e:
            if 'counter_frame' in locals(): counter_frame.destroy()
            messagebox.showerror("Erro", str(e))

    def tela_inicial(self):
        for w in self.root.winfo_children(): w.destroy()
        self.root.geometry("700x650")
        centralizar(self.root)

        main_container = ctk.CTkFrame(self.root, fg_color="transparent")
        main_container.pack(fill="both", expand=True, padx=40, pady=20)

        ctk.CTkLabel(main_container, text="MENU PRINCIPAL", font=FONT_TITLE, text_color="#2C3E50").pack(pady=20)

        btns_frame = ctk.CTkFrame(main_container, fg_color="transparent")
        btns_frame.pack(pady=10)

        def btn_style(btn):
            btn.configure(width=350, height=50, corner_radius=10, font=FONT_BUTTON)
            btn.pack(pady=8)

        # Botão Gerar Pendentes
        if not self._check_txt_today(TXT_ULTIMO_GERAR):
            btn_gerar = ctk.CTkButton(btns_frame, text="GERAR PENDENTES", fg_color="#2980B9", hover_color="#21618C")
            btn_style(btn_gerar)
            btn_gerar.configure(command=lambda b=btn_gerar: GerarPlanilhaWindow(self.root, self.usuario_atual, self, b))

        # Botão Salvar Planilhas
        if not self._check_txt_today(TXT_ULTIMO_ENVIO):
            btn_salvar = ctk.CTkButton(btns_frame, text="SALVAR PLANILHAS", fg_color="#27AE60", hover_color="#1E8449")
            btn_style(btn_salvar)
            btn_salvar.configure(command=lambda b=btn_salvar: threading.Thread(target=self.salvar_planilhas_outlook, args=(b, self.usuario_atual), daemon=True).start())

        # Botão Ajustar Planilhas
        if not self._check_txt_today(TXT_ULTIMO_AJUSTE):
            btn_ajustar = ctk.CTkButton(btns_frame, text="AJUSTAR PLANILHAS", fg_color="#F39C12", hover_color="#D68910")
            btn_style(btn_ajustar)
            btn_ajustar.configure(command=lambda b=btn_ajustar: threading.Thread(target=self._ajustar_planilhas_thread, args=(b,), daemon=True).start())

        # Envio / Cobrança
        btn_ec = ctk.CTkButton(btns_frame, text="ENVIO / COBRANÇA", fg_color="#8E44AD", hover_color="#71368A", command=self.tela_envio_cobranca)
        btn_style(btn_ec)

        # Procurações
        btn_proc = ctk.CTkButton(btns_frame, text="PROCURAÇÕES", fg_color="#34495E", hover_color="#283747", command=self.tela_procuracoes)
        btn_style(btn_proc)

        self.status_frame = ctk.CTkFrame(main_container, fg_color="transparent")
        self.status_frame.pack(fill="x", pady=20)

        rodape = ctk.CTkFrame(self.root, height=50, fg_color="transparent")
        rodape.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=10)
        ctk.CTkButton(rodape, text="SAIR", font=FONT_BUTTON, width=100, fg_color="#E74C3C", command=self.confirmar_saida).pack(side=tk.LEFT)
        help_link = ctk.CTkLabel(rodape, text="ENCONTROU UM ERRO? CLIQUE AQUI", font=FONT_FORGOT_PASSWORD, text_color="#3498DB", cursor="hand2")
        help_link.pack(side=tk.RIGHT)
        help_link.bind("<Button-1>", abrir_link_ajuda)

    def _check_txt_today(self, filename):
        if os.path.isfile(filename):
            with open(filename, "r") as f:
                d = f.read().strip()
                try:
                    dt = datetime.strptime(d, "%d/%m/%Y")
                    return dt.month == datetime.now().month and dt.year == datetime.now().year
                except: pass
        return False

    def confirmar_saida(self):
        if messagebox.askyesno("Sair", f"{self.usuario_atual.upper()}, deseja realmente sair?"):
            self.tela_login()

    def tela_envio_cobranca(self):
        has_envio = self._has_envio_actions()
        has_cobranca = self._has_cobranca_actions()
        if not has_envio and not has_cobranca:
            messagebox.showinfo("Aviso", "Não há ações disponíveis.")
            return

        win = ctk.CTkToplevel(self.root)
        win.title("E.A.I. - ENVIO / COBRANÇA")
        win.geometry("500x400")
        centralizar_janela(win, 500, 400)

        ctk.CTkLabel(win, text="ESCOLHA A AÇÃO", font=FONT_TITLE).pack(pady=30)

        if has_envio:
            ctk.CTkButton(win, text="ENVIO", font=FONT_BUTTON, width=300, height=50, fg_color="#2980B9", command=self.tela_empresas_envio).pack(pady=10)
        if has_cobranca:
            ctk.CTkButton(win, text="COBRANÇA", font=FONT_BUTTON, width=300, height=50, fg_color="#8E44AD", command=self.tela_empresas_cobranca).pack(pady=10)

        ctk.CTkButton(win, text="VOLTAR", font=FONT_BUTTON, fg_color="#E74C3C", command=win.destroy).pack(pady=30)

    def tela_empresas_envio(self):
        win = ctk.CTkToplevel(self.root)
        win.title("E.A.I. - ENVIO LISTAGEM")
        win.after(200, lambda: win.state('zoomed'))

        frame_top = ctk.CTkFrame(win, fg_color="transparent")
        frame_top.pack(side=tk.TOP, fill=tk.X, pady=20)
        ctk.CTkLabel(frame_top, text="ENVIO DE LISTAGEM", font=FONT_TITLE).pack()

        ctk.CTkButton(frame_top, text="ENVIAR TODOS", font=FONT_BUTTON, fg_color="#1F5FBF", command=lambda: threading.Thread(target=self._enviar_todos_thread, daemon=True).start()).pack(pady=10)

        registros = carregar_txt(TXT_ENVIO)
        botoes = [lbl for lbl in self.labels if not self._enviado_no_mes(lbl, registros) and not self._has_ok_status(lbl)]
        botoes.sort(key=lambda x: normalize_key(x))

        self.contador_label_envio = ctk.CTkLabel(frame_top, text=f"PENDENTES ESTE MÊS: {len(botoes)}", font=FONT_LABEL)
        self.contador_label_envio.pack()

        scroll = ctk.CTkScrollableFrame(win, orientation="horizontal")
        scroll.pack(fill="both", expand=True, padx=20, pady=10)

        inner = ctk.CTkFrame(scroll, fg_color="transparent")
        inner.pack()
        self.frame_empresas_envio = inner

        for i, label in enumerate(botoes):
            r, c = i % MAX_LINHAS_COLUNA, i // MAX_LINHAS_COLUNA
            btn = ctk.CTkButton(inner, text=label, font=FONT_BUTTON, width=300, height=80, corner_radius=10)
            btn.configure(command=lambda l=label, b=btn: threading.Thread(target=self._enviar_uma_label_thread, args=(l, b), daemon=True).start())
            btn.grid(row=r, column=c, padx=10, pady=10)

        ctk.CTkButton(win, text="VOLTAR", font=FONT_BUTTON, fg_color="#E74C3C", command=win.destroy).pack(pady=10)

    def _enviado_no_mes(self, label, registros=None):
        if registros is None: registros = carregar_txt(TXT_ENVIO)
        if label not in registros: return False
        try:
            dt = datetime.strptime(registros[label], "%d/%m/%Y")
            return dt.month == datetime.now().month and dt.year == datetime.now().year
        except: return False

    def _enviar_uma_label_thread(self, label, btn):
        success, err = enviar_email_label(label, self.df, self.usuario_atual)
        def ui():
            if success:
                try: btn.destroy()
                except: pass
            else: messagebox.showerror("Erro", f"Falha ao enviar '{label}': {err}")
            if self.frame_empresas_envio and self.frame_empresas_envio.winfo_exists():
                btns = [w for w in self.frame_empresas_envio.winfo_children() if isinstance(w, ctk.CTkButton)]
                self.contador_label_envio.configure(text=f"PENDENTES ESTE MÊS: {len(btns)}")
        self.root.after(0, ui)

    def _enviar_todos_thread(self):
        if not self.frame_empresas_envio: return
        widgets = [w for w in self.frame_empresas_envio.winfo_children() if isinstance(w, ctk.CTkButton)]
        for w in widgets:
            self._enviar_uma_label_thread(w.cget("text"), w)
            time.sleep(0.5)

    def _has_ok_status(self, label):
        row = get_row_by_label(self.df, label)
        if row is None: return False
        return str(row.iloc[IDX_AA_STATUS]).strip().upper() == "OK"

    def tela_empresas_cobranca(self):
        win = ctk.CTkToplevel(self.root)
        win.title("E.A.I. - COBRANÇA LISTAGEM")
        win.after(200, lambda: win.state('zoomed'))

        frame_top = ctk.CTkFrame(win, fg_color="transparent")
        frame_top.pack(side=tk.TOP, fill=tk.X, pady=20)
        ctk.CTkLabel(frame_top, text="COBRANÇA DE LISTAGEM", font=FONT_TITLE).pack()

        ctk.CTkButton(frame_top, text="COBRAR TODOS", font=FONT_BUTTON, fg_color="#1F5FBF", command=lambda: threading.Thread(target=self._cobrar_todos_thread, daemon=True).start()).pack(pady=10)

        registros = carregar_txt(TXT_COBRANCA)
        botoes = [lbl for lbl in self.labels if not self._cobrado_hoje(lbl, registros) and lbl not in EXCLUIR_NA_COBRANCA and not self._has_ok_status(lbl)]
        botoes.sort(key=lambda x: normalize_key(x))

        self.contador_label_cobranca = ctk.CTkLabel(frame_top, text=f"PENDENTES HOJE: {len(botoes)}", font=FONT_LABEL)
        self.contador_label_cobranca.pack()

        scroll = ctk.CTkScrollableFrame(win, orientation="horizontal")
        scroll.pack(fill="both", expand=True, padx=20, pady=10)

        inner = ctk.CTkFrame(scroll, fg_color="transparent")
        inner.pack()
        self.frame_empresas_cobranca = inner

        for i, label in enumerate(botoes):
            r, c = i % MAX_LINHAS_COLUNA, i // MAX_LINHAS_COLUNA
            btn = ctk.CTkButton(inner, text=label, font=FONT_BUTTON, width=300, height=80, corner_radius=10)
            btn.configure(command=lambda l=label, b=btn: threading.Thread(target=self._cobrar_uma_label_thread, args=(l, b), daemon=True).start())
            btn.grid(row=r, column=c, padx=10, pady=10)

        ctk.CTkButton(win, text="VOLTAR", font=FONT_BUTTON, fg_color="#E74C3C", command=win.destroy).pack(pady=10)

    def _cobrado_hoje(self, label, registros=None):
        if registros is None: registros = carregar_txt(TXT_COBRANCA)
        if label not in registros: return False
        try:
            dt = datetime.strptime(registros[label], "%d/%m/%Y")
            return dt.date() == datetime.now().date()
        except: return False

    def _cobrar_uma_label_thread(self, label, btn):
        success, err = cobrar_email_label(label, self.df, self.usuario_atual)
        def ui():
            if success:
                try: btn.destroy()
                except: pass
            else: messagebox.showerror("Erro", f"Falha ao cobrar '{label}': {err}")
            if self.frame_empresas_cobranca and self.frame_empresas_cobranca.winfo_exists():
                btns = [w for w in self.frame_empresas_cobranca.winfo_children() if isinstance(w, ctk.CTkButton)]
                self.contador_label_cobranca.configure(text=f"PENDENTES HOJE: {len(btns)}")
        self.root.after(0, ui)

    def _cobrar_todos_thread(self):
        if not self.frame_empresas_cobranca: return
        widgets = [w for w in self.frame_empresas_cobranca.winfo_children() if isinstance(w, ctk.CTkButton)]
        for w in widgets:
            self._cobrar_uma_label_thread(w.cget("text"), w)
            time.sleep(0.5)

def get_row_by_label(df, label):
    mask = df.iloc[:, IDX_AK_LABEL].astype(str) == label
    return df[mask].iloc[0] if mask.any() else None

def get_row_by_name(df, name):
    mask = df.iloc[:, 0].astype(str).str.upper() == name.upper()
    return df[mask].iloc[0] if mask.any() else None

def enviar_email_procuracao(empresa, df, usuario):
    try:
        pythoncom.CoInitialize()
        outlook = win32.Dispatch("Outlook.Application")
        session = outlook.Session
        row = get_row_by_name(df, empresa)
        if row is None: raise Exception(f"Empresa '{empresa}' não encontrada.")
        email_to = safe_email_field(row.iloc[IDX_AE_PROCURACAO_TO])
        email_cc = safe_email_field(row.iloc[IDX_AF_PROCURACAO_CC])
        if not email_to: raise Exception(f"E-mail não configurado.")
        mail = outlook.CreateItem(0)
        email_conta = EMAIL_CONTA_PADRAO_THAILANY if usuario == "thailany" else EMAIL_CONTA_PADRAO_VICTOR
        found = False
        for acc in session.Accounts:
            if str(acc.SmtpAddress).lower() == email_conta.lower():
                mail._oleobj_.Invoke(64209, 0, 8, 0, acc); found = True; break
        if not found: raise Exception(f"Conta {email_conta} não encontrada.")
        mail.To, mail.CC, mail.Subject = email_to, email_cc, f"Procuração eletrônica - {empresa}"
        mail.Display(False); time.sleep(0.25)
        corpo_html = f"""
<html><body><div style='font-family: Calibri; font-size: 11pt;'>
<p>Prezados,</p><p>A procuração eletrônica para acesso ao E-Social está vencida. Solicitamos a renovação.</p>
<p>Instruções no site do <a href="https://cav.receita.fazenda.gov.br/autenticacao/login">E-CAC</a>:</p>
<ol><li>Acessar E-CAC;</li><li>Procuração Eletrônica – Cadastra Procuração;</li><li>CPF/CNPJ Procurador – <b>18.802.823/0001-04;</b></li>
<li>Selecionar E-Social e Cadastrar (SST);</li><li>Assinar documento.</li></ol>
<p><b>Vigência de 2 ou mais anos.</b></p>
<p><span style='background-color: yellow;'><b>CNPJ OSWALDO CRUZ: 18.802.823/0001-04</b></span></p>
</div></body></html>"""
        mail.HTMLBody = corpo_html + (mail.HTMLBody or "")
        mail.Send(); registrar_txt(TXT_PROCURACOES, empresa)
        return True, None
    except Exception as e: return False, str(e)
    finally: pythoncom.CoUninitialize()

def get_anexos_from_row(row, usuario):
    try: raw = str(row.iloc[IDX_AL_ANEXOS])
    except: raw = ""
    if not raw or raw.lower() in {"nan", "none"}: return []
    anexos = []
    for p in re.split(r"[;,]", raw):
        p_c = limpar_caminho(p)
        if p_c and os.path.isfile(p_c): anexos.append(p_c)
    return anexos

def get_assunto_from_row(row):
    try: return safe_email_field(str(row.iloc[IDX_AM_ASSUNTO]))
    except: return ""

def listar_labels(df):
    labels = list(df.iloc[2:, IDX_AK_LABEL].dropna().astype(str))
    labels.sort(key=lambda x: normalize_key(x))
    return labels

def abrir_link_ajuda(event=None):
    webbrowser.open("https://wa.me/551221239207?text=Encontrei%20um%20erro%20no%20sistema%20E.A.I.%20Pode%20me%20ajudar%3F")

if __name__ == "__main__":
    # Suporte a DPI para Windows
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
    except: pass

    root = ctk.CTk()
    root.withdraw()
    app = App(root)
    def abrir_login():
        root.deiconify()
        app.usuario_entry.focus_set()
    mostrar_intro(root, abrir_login)
    root.mainloop()
